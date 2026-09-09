"""Persistent workflow, quality gate and publication regression coverage."""
import copy
import time

import pytest
from fastapi.testclient import TestClient

from semantica.explorer.app import create_app
from semantica.explorer.pipelines.definitions import EXTRACTION, FLOW, RULES, validate_definition
from semantica.explorer.pipelines.handlers import validate_candidate
from semantica.explorer.pipelines.repository import Conflict


def test_docx_tables_textboxes_and_media_coverage(tmp_path):
    import zipfile
    from semantica.explorer.pipelines.handlers import parse_document, chunk_document
    path = tmp_path / 'structured.docx'
    xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:body>
    <w:p><w:r><w:t>甲医院与乙学院合作。</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>疾病</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>营养建议</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    <w:p><mc:AlternateContent><mc:Choice><w:txbxContent><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p></w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent><w:p><w:r><w:t>重复回退文本</w:t></w:r></w:p></w:txbxContent></mc:Fallback></mc:AlternateContent></w:p>
    </w:body></w:document>'''
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', xml)
        archive.writestr('word/media/image1.png', b'fixture')
    document = parse_document(path, path.name)
    assert [s['text'] for s in document['segments']] == ['甲医院与乙学院合作。', '疾病 | 营养建议', '文本框内容']
    assert len(chunk_document(document, FLOW)['chunks']) == 1
    value, _ = extraction(None, None)
    from semantica.explorer.pipelines.handlers import normalize
    candidate = validate_candidate(normalize([value], RULES), document, EXTRACTION, RULES)
    assert any(i['rule'].startswith('document_coverage:') for i in candidate['issues'])


def test_docx_numeric_page_footer_does_not_trigger_coverage_warning(tmp_path):
    """A PAGE field in a footer is layout metadata, not missed source text."""
    import zipfile
    from semantica.explorer.pipelines.handlers import parse_document

    path = tmp_path / 'page-number-footer.docx'
    document_xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:body><w:p><w:r><w:t>正文事实</w:t></w:r></w:p></w:body></w:document>'''
    footer_xml = '''<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE \\* MERGEFORMAT </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>6</w:t></w:r></w:p></w:ftr>'''
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', document_xml)
        archive.writestr('word/footer1.xml', footer_xml)
    document = parse_document(path, path.name)
    assert document['coverage_issues'] == []


def test_xlsx_rows_headers_types_and_sheet_boundaries(tmp_path):
    from datetime import datetime
    from openpyxl import Workbook
    from semantica.explorer.pipelines.handlers import parse_document, chunk_document
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '机构'
    sheet.append(['编号', '机构', '人数', '日期'])
    sheet.append([12, '甲医院', 0, datetime(2026, 9, 8)])
    sheet['A2'].number_format = '00000'
    sheet.append(['00013', '乙学院', None, None])
    other = workbook.create_sheet('合作')
    other.append(['主体', '客体'])
    other.append(['甲医院', '乙学院'])
    path = tmp_path / 'rows.xlsx'
    workbook.save(path)
    document = parse_document(path, path.name)
    assert not document['coverage_issues']
    row = document['segments'][1]
    assert row['sheet'] == '机构' and row['row'] == 2
    assert '编号 (A2)=00012' in row['text']
    assert '人数 (C2)=0' in row['text']
    assert '2026-09-08' in row['text']
    assert '00013' in document['segments'][2]['text']
    chunks = chunk_document(document, {'chunk_size': 180, 'overlap': 0})['chunks']
    assert all(len({s['sheet'] for s in c['segments']}) == 1 for c in chunks)
    assert all('机构 (B' in s['text'] for s in document['segments'][1:3])


def test_xlsx_formula_and_merged_cells_enter_quality_gate(tmp_path):
    from openpyxl import Workbook
    from semantica.explorer.pipelines.handlers import parse_document
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['名称', '合计'])
    sheet.append(['甲医院', '=1+2'])
    sheet.merge_cells('A3:B3')
    sheet['A3'] = '合并说明'
    path = tmp_path / 'formula.xlsx'
    workbook.save(path)
    document = parse_document(path, path.name)
    assert any('公式' in issue for issue in document['coverage_issues'])
    assert any('合并' in issue for issue in document['coverage_issues'])
    assert '=1+2' not in str(document['segments'])


def test_xlsx_upload_enters_pipeline(client):
    import io
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['机构', '合作机构'])
    workbook.active.append(['甲医院', '乙学院'])
    stream = io.BytesIO()
    workbook.save(stream)
    assert '.xlsx' in client.get('/api/pipeline-capabilities').json()['formats']
    upload = client.post('/api/pipeline-documents', files={'file': ('hospital.xlsx', stream.getvalue())},
                         data={'request_key': 'xlsx-upload'})
    assert upload.status_code == 201
    body = payload(client)
    body['document_id'] = upload.json()['id']
    response = client.post('/api/pipeline-runs', json=body)
    assert response.status_code == 202
    run = wait(client, response.json()['id'])
    assert all(step['status'] == 'completed' for step in run['steps'])
    assert run['status'] == 'awaiting_review'


def test_xlsx_empty_and_oversized_dimensions_rejected(tmp_path):
    from openpyxl import Workbook
    from semantica.explorer.pipelines.handlers import parse_document
    workbook = Workbook()
    path = tmp_path / 'empty.xlsx'
    workbook.save(path)
    with pytest.raises(ValueError, match='no readable text'):
        parse_document(path, path.name)
    workbook.active.cell(1, 257, 'too wide')
    workbook.save(path)
    with pytest.raises(ValueError, match='256-column'):
        parse_document(path, path.name)


def extraction(chunk, template):
    return {'entities': [{'name': '甲医院', 'type': '机构'}, {'name': '乙学院', 'type': '机构'}],
            'relationships': [{'source': '甲医院', 'target': '乙学院', 'type': '合作', 'status': '已运营',
                               'segment_ids': ['p1'], 'quantity_text': None, 'year_text': None}]}, {'total_tokens': 12}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('SEMANTICA_PIPELINE_ENABLED', 'true')
    monkeypatch.setenv('SEMANTICA_PIPELINE_DATA_DIR', str(tmp_path))
    with TestClient(create_app()) as c:
        c.app.state.pipeline_service.model_call = extraction
        yield c


def payload(client, policy='manual'):
    doc = client.post('/api/pipeline-documents', files={'file': ('hospital.txt', '甲医院与乙学院合作。'.encode())},
                      data={'request_key': 'upload-fixture-1'})
    assert doc.status_code == 201, doc.text
    definitions = client.get('/api/pipeline-definitions').json()['items']
    return {'document_id': doc.json()['id'], 'model_profile_id': 'deepseek', 'publish_policy': policy,
            'request_key': 'run-fixture-key', **{d['kind'] + '_version_id': d['versions'][0]['id'] for d in definitions}}


def wait(client, identity):
    for _ in range(300):
        run = client.get('/api/pipeline-runs/' + identity).json()
        if run['status'] not in ('queued', 'running', 'publishing'):
            return run
        time.sleep(.03)
    pytest.fail('Run timed out')


def test_real_engine_candidate_and_event_history(client):
    body = payload(client)
    response = client.post('/api/pipeline-runs', json=body)
    assert response.status_code == 202, response.text
    identity = response.json()['id']
    assert client.post('/api/pipeline-runs', json=body).json()['id'] == identity
    run = wait(client, identity)
    assert run['status'] == 'awaiting_review', run
    assert all(s['status'] == 'completed' for s in run['steps'])
    assert len(client.app.state.session.graph.nodes) == 0
    result = client.get('/api/pipeline-runs/' + identity + '/results').json()['candidate']
    assert result['relationships'][0]['status'] == '文档陈述'
    assert result['audit'][0]['rule'] == 'cooperation_not_operation'
    events = client.get('/api/pipeline-runs/' + identity + '/events?limit=500').json()
    assert events['next_seq'] == run['last_event_seq']
    assert len([e for e in events['items'] if e['event'] == 'step.started']) == 6
    assert len([e for e in events['items'] if e['event'] == 'step.completed']) == 6


def test_neo4j_target_waits_for_alignment_even_with_auto_policy(client, monkeypatch):
    monkeypatch.setenv('NEO4J_URI', 'bolt://localhost:17690')
    monkeypatch.setenv('NEO4J_PASSWORD', 'test-only')
    body = {**payload(client, 'auto_if_clean'), 'publish_target': 'neo4j'}
    response = client.post('/api/pipeline-runs', json=body)
    assert response.status_code == 202, response.text
    run = wait(client, response.json()['id'])
    assert run['status'] == 'awaiting_review'
    assert run['quality'] == 'passed'
    assert run['publish_target'] == 'neo4j'
    assert run['publication'] is None


def test_review_can_restore_missing_entity_with_audit_and_revalidation(client):
    body = payload(client)
    def missing_endpoint(chunk, template):
        value, usage = extraction(chunk, template)
        value['entities'] = value['entities'][:1]
        return value, usage
    client.app.state.pipeline_service.model_call = missing_endpoint
    identity = client.post('/api/pipeline-runs', json=body).json()['id']
    run = wait(client, identity)
    endpoint = '/api/pipeline-runs/' + identity
    assert client.get(endpoint + '/results').json()['candidate']['issues']
    change = {'request_key': 'restore-entity', 'expected_candidate_revision': run['candidate_revision'],
              'reason': '原文明确提及乙学院，补回遗漏的关系端点',
              'additions': {'entities': [{'name': '乙学院', 'type': '机构'}]}}
    response = client.put(endpoint + '/review', json=change)
    assert response.status_code == 200, response.text
    candidate = client.get(endpoint + '/results').json()['candidate']
    assert candidate['issues'] == []
    assert any(a['rule'] == 'human_add' and a['reason'] == change['reason'] for a in candidate['audit'])
    assert client.put(endpoint + '/review', json=change).json() == response.json()
    invalid = {**change, 'request_key': 'invalid-entity', 'expected_candidate_revision': response.json()['candidate_revision'],
               'additions': {'entities': [{'name': '原文没有的机构', 'type': '机构'}]}}
    assert client.put(endpoint + '/review', json=invalid).status_code == 200
    assert any(i['rule'] == 'entity_source' for i in client.get(endpoint + '/results').json()['candidate']['issues'])


def test_auto_publication_is_idempotent_and_does_not_mutate_live_graph(client):
    body = payload(client, 'auto_if_clean')
    identity = client.post('/api/pipeline-runs', json=body).json()['id']
    run = wait(client, identity)
    assert run['status'] == 'completed', run
    assert run['publication']['nodes'] == 2
    assert len(client.app.state.session.graph.nodes) == 0
    publication = client.post('/api/pipeline-runs/' + identity + '/publish', json={'candidate_revision': 1, 'request_key': 'publish-again'}).json()
    assert publication['id'] == run['publication']['id']
    graph = client.get('/api/pipeline-graphs/' + publication['id'] + '/download').json()
    assert len(graph['nodes']) == 2 and len(graph['edges']) == 1
    response = client.post('/api/pipeline-graphs/' + publication['id'] + '/activate', json={'revision': 1, 'confirm_switch': True})
    assert response.status_code == 200, response.text
    assert len(client.app.state.session.graph.nodes) == 2


def test_auto_publish_failure_keeps_candidate_retryable(client, monkeypatch):
    from semantica.context import ContextGraph
    original = ContextGraph.save_to_file
    def fail(*args, **kwargs):
        raise OSError('injected disk failure')
    monkeypatch.setattr(ContextGraph, 'save_to_file', fail)
    body = payload(client, 'auto_if_clean')
    identity = client.post('/api/pipeline-runs', json=body).json()['id']
    run = wait(client, identity)
    for _ in range(100):
        run = client.get('/api/pipeline-runs/' + identity).json()
        if run.get('error'):
            break
        time.sleep(.03)
    assert run['status'] == 'awaiting_review'
    assert run['error'] and not run['publication']
    monkeypatch.setattr(ContextGraph, 'save_to_file', original)
    response = client.post('/api/pipeline-runs/' + identity + '/publish',
                           json={'candidate_revision': 1, 'request_key': 'manual-retry'})
    assert response.status_code == 200, response.text


def test_stale_queue_claim_cannot_restart_cancelled_task(client):
    svc = client.app.state.pipeline_service
    with svc.repo.transaction():
        svc.repo.put('run', 'cancelled-fixture', {'id': 'cancelled-fixture', 'status': 'cancelled', 'last_event_seq': 0})
    svc._execute('cancelled-fixture')
    assert svc.repo.get('run', 'cancelled-fixture')['status'] == 'cancelled'
    assert svc.repo.events('cancelled-fixture', 0, 100)['items'] == []


def test_startup_marks_uncommitted_work_interrupted(tmp_path):
    from semantica.explorer.pipelines.repository import Repository
    from semantica.explorer.pipelines.service import PipelineService
    repo = Repository(tmp_path)
    with repo.transaction():
        repo.put('run', 'interrupted-fixture', {'id': 'interrupted-fixture', 'status': 'running',
                 'last_event_seq': 0, 'steps': [{'id': 'parse', 'status': 'running'}, {'id': 'split', 'status': 'pending'}]})
    repo.close()
    svc = PipelineService(tmp_path, model_call=extraction)
    try:
        run = svc.repo.get('run', 'interrupted-fixture')
        assert run['status'] == 'interrupted'
        assert all(s['status'] == 'blocked' for s in run['steps'])
        assert svc.repo.events(run['id'], 0, 10)['items'][0]['event'] == 'run.recovered'
    finally:
        svc.close()


def test_draft_isolation_and_preview_publish_rejection(client):
    body = payload(client)
    definition = client.get('/api/pipeline-definitions/default-extraction').json()
    updated = {**definition['body'], 'instructions': 'Changed'}
    saved = client.put('/api/pipeline-definitions/default-extraction', json={'expected_revision': 1, 'body': updated})
    assert saved.status_code == 200
    assert client.put('/api/pipeline-definitions/default-extraction', json={'expected_revision': 1, 'body': updated}).status_code == 409
    old = client.get('/api/pipeline-definition-versions/' + body['extraction_version_id']).json()
    assert old['body']['instructions'] != 'Changed'
    body['extraction_version_id'] = 'draft:default-extraction:2'
    identity = client.post('/api/pipeline-previews', json=body).json()['id']
    assert wait(client, identity)['status'] == 'preview_completed'
    assert client.post('/api/pipeline-runs/' + identity + '/publish', json={'candidate_revision': 1, 'request_key': 'preview-publish'}).status_code == 409


def test_rule_rerun_reuses_model_artifact(client):
    calls = []
    client.app.state.pipeline_service.model_call = lambda *args: (calls.append(1) or extraction(*args))
    body = payload(client)
    identity = client.post('/api/pipeline-runs', json=body).json()['id']
    wait(client, identity)
    definition = client.get('/api/pipeline-definitions/default-rules').json()
    new_rules = {**definition['body'], 'required_mentions': []}
    client.put('/api/pipeline-definitions/default-rules', json={'expected_revision': 1, 'body': new_rules})
    version = client.post('/api/pipeline-definitions/default-rules/versions', json={'expected_revision': 2, 'request_key': 'version-key'}).json()
    body.update(rules_version_id=version['id'], request_key='new-run-key')
    plan = client.post('/api/pipeline-runs/' + identity + '/rerun-plan', json=body).json()
    assert set(plan['reuse']) == {'parse', 'split', 'extract'}
    response = client.post('/api/pipeline-runs/' + identity + '/reruns', json={**body, 'plan_hash': plan['hash']})
    assert response.status_code == 202, response.text
    child = wait(client, response.json()['id'])
    assert child['parent_run_id'] == identity and len(calls) == 1
    assert child['steps'][2]['status'] == 'reused'
    comparison = client.get('/api/pipeline-runs/' + child['id'] + '/comparison', params={'other_run_id': identity}).json()
    assert set(comparison['configuration']) == {'rules'}
    assert comparison['quality'] == {'before': [], 'after': []}


def test_failed_extraction_blocks_downstream(client):
    def fail(*args):
        raise ValueError('MODEL_INCOMPLETE')
    client.app.state.pipeline_service.model_call = fail
    identity = client.post('/api/pipeline-runs', json=payload(client)).json()['id']
    run = wait(client, identity)
    assert run['status'] == 'failed'
    assert run['steps'][2]['status'] == 'failed'
    assert all(s['status'] == 'blocked' for s in run['steps'][3:])
    assert not run['publication']


def test_review_cas_and_revalidate(client):
    body = payload(client)
    identity = client.post('/api/pipeline-runs', json=body).json()['id']
    wait(client, identity)
    record = client.get('/api/pipeline-runs/' + identity + '/results').json()['candidate']['relationships'][0]
    patch = {'expected_candidate_revision': 1, 'request_key': 'review-request', 'reason': 'test evidence',
             'patches': [{'record_id': record['id'], 'changes': {'year_text': '2030年'}}]}
    updated = client.put('/api/pipeline-runs/' + identity + '/review', json=patch)
    assert updated.status_code == 200 and updated.json()['quality'] == 'needs_review'
    assert client.put('/api/pipeline-runs/' + identity + '/review', json={**patch, 'request_key': 'review-stale'}).status_code == 409
    assert client.post('/api/pipeline-runs/' + identity + '/publish', json={'candidate_revision': 2, 'request_key': 'publish-invalid'}).status_code == 409


def test_idempotency_conflict_and_auth(client, monkeypatch):
    body = payload(client)
    client.post('/api/pipeline-runs', json=body)
    assert client.post('/api/pipeline-runs', json={**body, 'publish_policy': 'auto_if_clean'}).status_code == 409
    monkeypatch.setenv('SEMANTICA_ALLOW_ANONYMOUS', 'false')
    monkeypatch.setenv('SEMANTICA_API_KEY', 'test-key')
    assert client.get('/api/pipeline-runs').status_code == 401
    assert client.get('/api/pipeline-runs', headers={'X-API-Key': 'test-key'}).status_code == 200


def test_flow_cannot_skip_handlers_or_create_branch():
    with pytest.raises(ValueError):
        validate_definition('flow', {**FLOW, 'steps': ['parse', 'extract', 'store']})
    with pytest.raises(ValueError):
        validate_definition('flow', {**FLOW, 'steps': FLOW['steps'] + ['execute_python']})


def test_subject_local_operational_evidence():
    document = {'segments': [{'id': 'p1', 'text': '甲医院未运营，乙学院已投入运营。', 'page': None, 'start': 0, 'end': 22}]}
    candidate = {'entities': [{'id': 'a', 'name': '甲医院', 'type': '机构'}, {'id': 'b', 'name': '乙学院', 'type': '设施类型'}],
                 'relationships': [{'id': 'r', 'source': '乙学院', 'target': '甲医院', 'type': '包含机构', 'status': '已运营', 'segment_ids': ['p1']}]}
    result = validate_candidate(candidate, document, EXTRACTION, RULES)
    assert any(i['rule'] == 'local_status_evidence' for i in result['issues'])


def test_artifact_corruption_is_detected(client):
    repo = client.app.state.pipeline_service.repo
    key = repo.artifact({'a': 1})
    (repo.root / 'artifacts' / key).write_text('{}')
    with pytest.raises(ValueError, match='checksum'):
        repo.read(key)


def test_cancel_inflight_call_does_not_publish(client):
    import threading
    entered, release = threading.Event(), threading.Event()
    def blocking(*args):
        entered.set()
        release.wait(5)
        return extraction(*args)
    client.app.state.pipeline_service.model_call = blocking
    identity = client.post('/api/pipeline-runs', json=payload(client, 'auto_if_clean')).json()['id']
    try:
        assert entered.wait(5)
        response = client.post('/api/pipeline-runs/' + identity + '/cancel', json={'request_key': 'cancel-request'})
        assert response.json()['status'] == 'cancel_requested'
    finally:
        release.set()
    for _ in range(100):
        run = client.get('/api/pipeline-runs/' + identity).json()
        if run['status'] == 'cancelled':
            break
        time.sleep(.03)
    assert run['status'] == 'cancelled'
    assert not run['publication']


def test_failed_publication_has_no_partial_commit(client, monkeypatch):
    from semantica.context import ContextGraph
    identity = client.post('/api/pipeline-runs', json=payload(client)).json()['id']
    wait(client, identity)
    original = ContextGraph.save_to_file
    def fail(*args, **kwargs):
        raise ValueError('Injected disk failure')
    monkeypatch.setattr(ContextGraph, 'save_to_file', fail)
    body = {'candidate_revision': 1, 'request_key': 'retry-same-publication'}
    assert client.post('/api/pipeline-runs/' + identity + '/publish', json=body).status_code == 422
    assert client.get('/api/pipeline-graphs').json()['items'] == []
    assert len(client.app.state.session.graph.nodes) == 0
    monkeypatch.setattr(ContextGraph, 'save_to_file', original)
    assert client.post('/api/pipeline-runs/' + identity + '/publish', json=body).status_code == 200


def test_completed_publication_survives_service_restart(tmp_path):
    from semantica.explorer.pipelines.service import PipelineService
    svc = PipelineService(tmp_path, model_call=extraction)
    try:
        doc = svc.upload('hospital.txt', '甲医院与乙学院合作。'.encode(), 'restart-upload')
        body = {'document_id': doc['id'], 'request_key': 'restart-run', 'publish_policy': 'auto_if_clean',
                **{d['kind'] + '_version_id': d['versions'][0]['id'] for d in svc.definitions()}}
        identity = svc.create_run(body)['id']
        for _ in range(200):
            run = svc.repo.get('run', identity)
            if run['status'] == 'completed':
                break
            time.sleep(.03)
        assert run['status'] == 'completed'
    finally:
        svc.close()
    restored = PipelineService(tmp_path, model_call=extraction)
    try:
        assert restored.repo.get('run', identity)['publication'] == run['publication']
        assert (tmp_path / 'artifacts' / run['publication']['artifact']).exists()
    finally:
        restored.close()


def test_template_overrides_do_not_mutate_published_template():
    from semantica.pipeline.pipeline_templates import PipelineTemplateManager
    manager = PipelineTemplateManager()
    before = copy.deepcopy(manager.get_template('document_processing').steps)
    manager.create_pipeline_from_template('document_processing', parse={'formats': ['txt']})
    assert manager.get_template('document_processing').steps == before


def test_engine_stop_is_not_reported_as_success():
    from semantica.pipeline import PipelineBuilder, ExecutionEngine
    from semantica.pipeline.execution_engine import PipelineStatus
    engine = ExecutionEngine(default_max_retries=0)
    builder = PipelineBuilder()
    builder.add_step('first', 'custom', handler=lambda data, **kw: engine.stop_pipeline('cancel-test'))
    called = []
    builder.add_step('second', 'custom', handler=lambda data, **kw: called.append(1))
    builder.connect_steps('first', 'second')
    result = engine.execute_pipeline(builder.build('cancel-test'))
    assert not result.success and not called
    assert engine.get_pipeline_status('cancel-test') == PipelineStatus.STOPPED
