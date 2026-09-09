"""Bounded, read-only workbook rows with repeated header context."""
import datetime
import re
import zipfile

from defusedxml import ElementTree


def read_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    issues = []
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('XLSX expands beyond 100 MiB')
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                with archive.open(name) as stream:
                    for _, node in ElementTree.iterparse(stream, events=('end',)):
                        if node.tag.endswith('}mergeCell'):
                            issues.append('工作簿包含合并单元格；当前按单元格原值读取，请审核多行表头及合并区域。')
                            break
                        node.clear()
        if any(name.startswith(('xl/media/', 'xl/charts/')) for name in archive.namelist()):
            issues.append('工作簿包含图片或图表，当前仅提取单元格内容。')

    def display(cell):
        value = cell.value
        if value is None:
            return None
        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
            return value.isoformat()
        if isinstance(value, bool):
            return 'true' if value else 'false'
        if isinstance(value, (int, float)) and re.fullmatch(r'0+', cell.number_format or ''):
            if value == int(value):
                return str(int(value)).zfill(len(cell.number_format))
        return str(value)

    paragraphs, characters, visited = [], 0, 0
    raw = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in raw:
                if sheet.max_column and sheet.max_column > 256:
                    raise ValueError('XLSX exceeds 256-column processing budget')
                if sheet.max_row and sheet.max_row > 100000:
                    raise ValueError('XLSX exceeds 100,000-row processing budget')
                if sheet.sheet_state != 'visible':
                    issues.append('工作表 %s 为隐藏状态，已读取，请核对其是否属于导入范围。' % sheet.title)
                headers = None
                for row_number, (row, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                    visited += len(row)
                    if visited > 1000000:
                        raise ValueError('XLSX exceeds 1,000,000-cell processing budget')
                    cells = []
                    for col, (cell, value) in enumerate(zip(row, values), 1):
                        text = display(value)
                        coordinate = '%s%s' % (get_column_letter(col), row_number)
                        if cell.data_type == 'f' and value.value is None:
                            issues.append('%s!%s 的公式没有已保存的计算结果，请在 Excel 中重新计算并保存。' % (sheet.title, coordinate))
                            text = '[公式结果缺失]'
                        if cell.data_type == 'e' or value.data_type == 'e':
                            issues.append('%s!%s 存在单元格错误，请修正后重试。' % (sheet.title, coordinate))
                        if text is not None:
                            cells.append((col, coordinate, text))
                    if not cells:
                        continue
                    if headers is None:
                        headers = {col: text for col, _, text in cells}
                        if any(not isinstance(c.value, str) or c.data_type == 'f' for c in row if c.value is not None):
                            issues.append('工作表 %s 首个非空行可能不是表头，请核对列含义。' % sheet.title)
                        body = ' | '.join('%s=%s' % (coord, text) for _, coord, text in cells)
                    else:
                        body = ' | '.join('%s (%s)=%s' % (headers.get(col, get_column_letter(col)), coord, text)
                                          for col, coord, text in cells)
                    text = '工作表：%s；第%s行；%s' % (sheet.title, row_number, body)
                    characters += len(text) + 2
                    if characters > 2000000:
                        raise ValueError('XLSX exceeds 2,000,000-character processing budget; split the workbook into smaller files')
                    paragraphs.append({'text': text, 'page': None, 'location': '%s!第%s行' % (sheet.title, row_number),
                                       'sheet': sheet.title, 'row': row_number,
                                       'headers': dict(headers or {}),
                                       'cells': {headers.get(col, get_column_letter(col)) if headers else get_column_letter(col): value
                                                 for col, _, value in cells}})
        finally:
            cached.close()
    finally:
        raw.close()
    return paragraphs, list(dict.fromkeys(issues))
