#!/usr/bin/env python3
"""Best-practice standalone importer: source registry, auditable chunks, curated KG, FAISS RAG."""
import argparse,hashlib,json,re,sys
from pathlib import Path
from datetime import datetime,timezone
from zipfile import ZipFile
from docx import Document
import openpyxl

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 ap=argparse.ArgumentParser();ap.add_argument('root');ap.add_argument('--out',default='medical_import_output/curated_hospital_knowledge');ap.add_argument('--no-embeddings',action='store_true');a=ap.parse_args(); root=Path(a.root);out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
 docs=[];chunks=[];tables=[]; issues=[]
 def add_chunks(text,meta,limit=900):
  text=re.sub(r'\s+',' ',text).strip()
  if not text:return
  for i in range(0,len(text),limit):
   x=text[i:i+limit].strip()
   if len(x)>=5: chunks.append({'chunk_id':f"{meta['document_id']}:{len(chunks):06d}",'text':x,'metadata':meta|{'offset':i}})
 for p in sorted(root.rglob('*')):
  if not p.is_file() or p.name=='.DS_Store' or p.suffix.lower() not in {'.docx','.xlsx'}:continue
  did='doc:'+hashlib.sha1(str(p.relative_to(root)).encode()).hexdigest()[:12]; rec={'document_id':did,'path':str(p),'relative_path':str(p.relative_to(root)),'format':p.suffix.lower()[1:],'sha256':sha(p),'bytes':p.stat().st_size};docs.append(rec)
  try:
   if p.suffix.lower()=='.docx':
    d=Document(p); para=0
    for para,x in enumerate(d.paragraphs,1):
     t=x.text.strip()
     if t:add_chunks(t,{'document_id':did,'source':str(p),'section':t if len(t)<100 else '正文','paragraph':para})
    for ti,table in enumerate(d.tables,1):
     for ri,row in enumerate(table.rows,1):
      vals=[c.text.strip() for c in row.cells]; add_chunks(' | '.join(vals),{'document_id':did,'source':str(p),'table':ti,'row':ri}); tables.append({'document_id':did,'source':str(p),'table':ti,'row':ri,'values':vals})
    rec['paragraphs']=len(d.paragraphs);rec['tables']=len(d.tables)
   else:
    w=openpyxl.load_workbook(p,data_only=True)
    rec['sheets']=w.sheetnames
    for ws in w.worksheets:
     rows=list(ws.values)
     if not rows:continue
     heads=[str(x).strip() if x is not None else '' for x in rows[0]]
     for rn,row in enumerate(rows[1:],2):
      vals={heads[i]:str(row[i]).strip() for i in range(min(len(heads),len(row))) if heads[i] and row[i] not in (None,'')}
      if vals:add_chunks('；'.join(f'{k}：{v}' for k,v in vals.items()),{'document_id':did,'source':str(p),'sheet':ws.title,'row':rn})
     rec.setdefault('sheet_stats',{})[ws.title]={'rows':ws.max_row,'columns':ws.max_column}
  except Exception as e:issues.append({'file':str(p),'error':repr(e)})
 (out/'documents.json').write_text(json.dumps(docs,ensure_ascii=False,indent=2),encoding='utf8')
 (out/'tables.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2),encoding='utf8')
 with (out/'chunks.jsonl').open('w',encoding='utf8') as f:
  for c in chunks:f.write(json.dumps(c,ensure_ascii=False)+'\n')
 # curated entity extraction: stable canonical IDs, evidence stays on entities/relations
 vocab={'营养不良':'营养状态','营养低下':'营养状态','老年营养':'营养主题','健康饮食':'营养主题','体重减轻':'症状','膳食摄入不足':'风险因素','恶性肿瘤':'疾病','炎症':'病理因素','肌肉量减少':'症状','低BMI':'指标','维生素D':'营养素','维生素B12':'营养素','蛋白质':'营养素','能量摄入':'营养指标','营养评估':'检查方法','GLIM':'诊断标准','吞咽困难':'症状','吸收不良':'病理因素','慢病':'专业方向','功能医学':'专业方向','精准医疗':'专业方向','抗衰老':'专业方向','心血管':'专业方向','免疫系统':'专业方向','健康管理':'专业方向','疾病超早期筛查':'专业方向','老年人':'人群'}
 es={};rs=[];rseen=set()
 def ent(t,n,ev=None):
  i=f'{t}:{n}';e=es.setdefault(i,{'id':i,'name':n,'type':t,'properties':{'evidence':[]}})
  if ev and ev not in e['properties']['evidence'] and len(e['properties']['evidence'])<50:e['properties']['evidence'].append(ev)
  return i
 h=ent('医院','循上鹏瑞利医院')
 for c in chunks:
  found=[(n,t) for n,t in vocab.items() if n in c['text']]
  ids=[ent(t,n,{'chunk_id':c['chunk_id'],'source':c['metadata'].get('source'),'location':{k:v for k,v in c['metadata'].items() if k in ('paragraph','row','sheet','table')}}) for n,t in found]
  for (n,t),i in zip(found,ids):
   if t=='专业方向':
    k=(h,i,'医院关注方向');
    if k not in rseen:rs.append({'source':h,'target':i,'type':'医院关注方向','evidence':[c['chunk_id']]});rseen.add(k)
  for j in range(len(ids)-1):
   k=(ids[j],ids[j+1],'同段相关');
   if k not in rseen:rs.append({'source':ids[j],'target':ids[j+1],'type':'同段相关','evidence':[c['chunk_id']]});rseen.add(k)
 # preserve structured doctor/group extraction from the Excel source, then merge by stable id
 prior=Path('medical_import_output/hospital_doctors_semantic/hospital_doctors_kg.json')
 if prior.exists():
   pk=json.loads(prior.read_text())
   for e in pk.get('entities',[]): es.setdefault(e['id'],e)
   for r in pk.get('relationships',[]):
    key=(r.get('source'),r.get('target'),r.get('type'))
    if key not in rseen: rs.append(r); rseen.add(key)
 kg={'entities':list(es.values()),'relationships':rs};(out/'knowledge_graph.json').write_text(json.dumps(kg,ensure_ascii=False),encoding='utf8');(out/'entities.json').write_text(json.dumps(kg['entities'],ensure_ascii=False,indent=2),encoding='utf8');(out/'relationships.json').write_text(json.dumps(kg['relationships'],ensure_ascii=False,indent=2),encoding='utf8')
 if not a.no_embeddings:
  from semantica.vector_store import VectorStore
  from semantica.embeddings import EmbeddingGenerator
  g=EmbeddingGenerator();g.set_text_model('fastembed','BAAI/bge-small-zh-v1.5');v=VectorStore(backend='faiss',dimension=512);v.embedder=g;v.add_documents([c['text'] for c in chunks],metadata=[c['metadata'] for c in chunks],batch_size=64);v.save(str(out/'rag_index'))
 report={'created_at':datetime.now(timezone.utc).isoformat(),'documents':len(docs),'chunks':len(chunks),'entities':len(es),'relationships':len(rs),'issues':issues,'entity_types':{t:sum(x['type']==t for x in es.values()) for t in set(x['type'] for x in es.values())},'rag':not a.no_embeddings};(out/'quality_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');(out/'import_manifest.json').write_text(json.dumps({'root':str(root),'documents':docs,'output':str(out)},ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
