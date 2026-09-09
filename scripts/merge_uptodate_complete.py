import json,sys,re
from pathlib import Path
from docx import Document
b=Path('medical_import_output/hospital_doctors_semantic'); p=b/'hospital_doctors_merged_kg.json'; k=json.loads(p.read_text()); es={e['id']:e for e in k['entities']}; rs=k['relationships']; seen={(r['source'],r['target'],r['type']) for r in rs}; fn=sys.argv[1]; d=Document(fn); h='医院:循上鹏瑞利医院'; current=None
for i,p0 in enumerate(d.paragraphs,1):
 t=re.sub(r'\s+',' ',p0.text).strip()
 if not t: continue
 typ='营养章节' if len(t)<100 and not t.startswith(('●','-')) else '营养知识段落'; ident=f'{typ}:uptodate:{i}'
 es[ident]={'id':ident,'name':t[:120],'type':typ,'properties':{'全文':t,'来源文件':fn,'段落':i,'样式':p0.style.name}}
 parent=current or h
 if typ=='营养章节': current=ident; parent=h
 key=(parent,ident,'包含内容')
 if key not in seen: rs.append({'source':parent,'target':ident,'type':'包含内容','properties':{'source':fn,'paragraph':i}});seen.add(key)
# output audit
p.write_text(json.dumps({'entities':list(es.values()),'relationships':rs},ensure_ascii=False),encoding='utf8')
print('paragraphs',sum(1 for x in d.paragraphs if x.text.strip()),'entities',len(es),'relationships',len(rs))
