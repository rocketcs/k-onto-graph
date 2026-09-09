import json,sys
from pathlib import Path
from docx import Document
from semantica.vector_store import VectorStore
from semantica.embeddings import EmbeddingGenerator
base=Path('medical_import_output/hospital_doctors_semantic'); xlsx=sys.argv[1]; docx=sys.argv[2]
import openpyxl
wb=openpyxl.load_workbook(xlsx,read_only=True,data_only=True); texts=[]; metas=[]
for ws in wb.worksheets:
 rows=list(ws.values)
 if not rows: continue
 heads=[str(x).strip() if x else '' for x in rows[0]]
 for rn,row in enumerate(rows[1:],2):
  v={heads[i]:str(row[i]).strip() for i in range(min(len(heads),len(row))) if heads[i] and row[i] not in (None,'')}
  if v.get('医生姓名') or v.get('特色医疗组'):
   texts.append('\n'.join(f'{k}：{val}' for k,val in v.items())); metas.append({'source':xlsx,'sheet':ws.title,'row':rn,'type':'医生介绍'})
d=Document(docx); paras=[p.text.strip() for p in d.paragraphs if p.text.strip()]
for i in range(0,len(paras),4):
 t='\n'.join(paras[i:i+4]);
 if len(t)>10: texts.append('医院介绍：\n'+t); metas.append({'source':docx,'section':'医院介绍','paragraph_start':i+1,'type':'医院介绍'})
g=EmbeddingGenerator();g.set_text_model('fastembed','BAAI/bge-small-zh-v1.5');v=VectorStore(backend='faiss',dimension=512);v.embedder=g;v.add_documents(texts,metadata=metas,batch_size=32); out=base/'hospital_doctors_merged_vector_index';v.save(str(out));print('chunks',len(texts),'index',out)
